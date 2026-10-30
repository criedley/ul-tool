#!/usr/bin/env python3
"""AP4 report tool: CTS max-temp stats + CTS charts, Agilent stats (cropped to CTS window).
Usage: python ap4_tool.py --cts CTS.xls --mapping map.xlsx [--agilent AG.csv] --out report.xlsx
"""
import re, argparse
from datetime import datetime, timedelta
import numpy as np, pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import PatternFill

RED = PatternFill("solid", start_color="FFFF0000", end_color="FFFF0000")

# ---------- CTS ----------
DATE_RE = re.compile(r"^\s*\d{4}-\d{2}-\d{2}")

CTS_TIME_FORMATS = ("%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S")

def _parse_cts_time(s):
    s = s.strip()
    for fmt in CTS_TIME_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None

def parse_cts(path):
    """CTS export: tab-delimited, ~12 lines of preamble, a header row whose first
    field is 'Date/Time', then units/AVE/MIN/MAX rows, then data rows."""
    header, rows, skipped = None, [], 0
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        for line in f:
            fields = line.rstrip("\r\n").split("\t")
            if header is None:
                if "Date/Time" in fields[0]:
                    while fields and not fields[-1].strip():
                        fields.pop()
                    header = [x.strip() for x in fields]
                continue
            ts = _parse_cts_time(fields[0])
            if ts is None:
                skipped += 1          # description / units / AVE / MIN / MAX / blank rows
                continue
            while len(fields) > len(header) and not fields[-1].strip():
                fields.pop()
            if len(fields) != len(header):
                raise ValueError(f"Data row at {ts} has {len(fields)} fields, header has {len(header)}")
            rows.append([ts] + fields[1:])
    if header is None:
        raise ValueError("No header line with 'Date/Time' as its first field was found")
    if len(rows) < 2:
        raise ValueError(f"Only {len(rows)} data rows parsed; check the date format")
    df = pd.DataFrame(rows, columns=header)
    for c in df.columns[1:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    print(f"CTS: {len(df)} rows, {len(header)} columns, {skipped} non-data lines skipped, "
          f"{df['Date/Time'].iloc[0]} .. {df['Date/Time'].iloc[-1]}")
    return df

# ---------- Agilent ----------
def parse_agilent(path):
    """Plain CSV: 4 metadata lines (Name/Owner/Comments/Acquisition Date), then a header
    row starting with 'Time', then 'm/d/yyyy HH:MM:SS,v1,...,vN' rows."""
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        skip = next((i for i, line in enumerate(f) if line.split(",", 1)[0].strip() == "Time"), None)
    if skip is None:
        raise ValueError("Agilent header row starting with 'Time' not found")
    df = pd.read_csv(path, skiprows=skip, encoding="utf-8-sig")
    df.columns = [c.strip() for c in df.columns]
    df["Time"] = pd.to_datetime(df["Time"], format="%m/%d/%Y %H:%M:%S")
    for c in df.columns[1:]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    print(f"Agilent: {len(df)} rows, {len(df.columns) - 1} channels, "
          f"{df['Time'].iloc[0]} .. {df['Time'].iloc[-1]}")
    return df

def trim_agilent(df, agstart, aglast, ti=None):
    """Keep Agilent rows inside the CTS window: first row >= CTS start, last row <= CTS end.
    Matches the cropped '_pertinent_' file (21:02:48 .. 01:02:28 for this test)."""
    out = df[(df["Time"] >= agstart) & (df["Time"] <= aglast)].reset_index(drop=True)
    if out.empty:
        raise ValueError(f"No Agilent rows between {agstart} and {aglast}; check clocks/dates")
    return out

# ---------- mapping ----------
def read_mapping(path):
    """Sheet1: [row#, Channel, Channel Name, Max Temps Allowed (C)] -> list of dicts."""
    ws = load_workbook(path, data_only=True).worksheets[0]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] is None:
            continue
        out.append({"channel": str(int(float(row[1]))),
                    "name": str(row[2] or "").strip(),
                    "limit_c": float(row[3]) if row[3] not in (None, "") else None})
    return out

# ---------- stats ----------
def f2c(f): return (f - 32) / 1.8

def calc(df, key, lim, lo, hi):
    if key not in df.columns: return None
    v = df[key].dropna(); used = v[(v >= lo) & (v <= hi)]
    if used.empty: return None
    s = np.sort(used.values)[::-1]
    mx = s[0]; top = s[4:10]
    avg = top.mean() if len(top) else np.nan
    maxc = f2c(mx)
    return dict(mx=mx, avg=avg, delta=abs(mx - avg), maxc=maxc,
                margin=None if lim is None else lim - maxc,
                used=len(used), excluded=len(v) - len(used))

def find_row(df, target, ti):
    """AP4.bas FindClosestRowtoTimeStringv2: scan up from the bottom; first row i where
    (t[i]-target <= ti+3 and t[i-1]-target >= -ti) or t[i]==target."""
    t = (df["Time"] - target).dt.total_seconds().values
    for i in range(len(t) - 1, 0, -1):
        if (t[i] <= ti + 3 and t[i - 1] >= -ti) or t[i] == 0:
            return i
    return None

def resolve_end(df, hhmm):
    """None/'0' -> last row of data. 'hh:mm' -> that time on whichever date falls inside the data."""
    first, last = df["Time"].iloc[0], df["Time"].iloc[-1]
    if not hhmm or hhmm == "0":
        return last
    h, m = map(int, hhmm.split(":"))
    for d in (last.date(), first.date()):
        cand = datetime(d.year, d.month, d.day, h, m)
        if first <= cand <= last:
            return cand
    raise ValueError(f"End time {hhmm} is outside the Agilent data ({first} .. {last})")

def write_stats(wb, title, df, items, lo, hi, end=None):
    ti = int(round(df["Time"].diff().dt.total_seconds().median()))
    end = end or df["Time"].iloc[-1]
    snap = [find_row(df, end - timedelta(minutes=m), ti) for m in (30, 15, 0)]   # "0min","15min","30min" columns
    print("Snapshot rows:", [str(df["Time"].iloc[i]) if i is not None else None for i in snap])
    ws = wb.create_sheet(title)
    ws.append(["Location", "Channel", "Key", "Max Temps Allowed (C)", "Max Temps (F) from data",
               "Average Top 5-10 temps (F)", "Delta", "Max Temps (C)", "Margin to limit (C)",
               "Samples used", "Samples out of range", "0min (C)", "15min (C)", "30min (C)", "Difference"])
    for loc, ch, key, lim in items:
        r = calc(df, key, lim, lo, hi)
        if r is None:
            ws.append([loc, ch, key, lim, "channel missing / no valid data"]); continue
        temps = []
        for i in snap:
            v = df[key].iloc[i] if i is not None else np.nan
            temps.append(f2c(v) if lo <= v <= hi else None)
        diff = max(temps) - min(temps) if None not in temps else None
        ws.append([loc, ch, key, lim, r["mx"], r["avg"], r["delta"], r["maxc"], r["margin"],
                   r["used"], r["excluded"], *temps, diff])
        row = ws.max_row
        if r["delta"] > 10: ws.cell(row, 7).fill = RED
        if r["margin"] is not None and r["margin"] < 0: ws.cell(row, 9).fill = RED
        if diff is not None and diff > 3: ws.cell(row, 15).fill = RED

# ---------- CTS charts (AP4Graphs.bas) ----------
# (sheet title, chart title, [(channel, series name, secondary?)], primary y title, secondary y title)
CHARTS = [
 ("ID Room Temps", "ID Room Temps", [("PID_PV_ID_DB_TEMP_STS", "ID Db Temps", 0), ("PID_PV_ID_WB_TEMP_STS", "ID Wb Temps", 1)], "Drybulb Temp (F)", "Wetbulb Temp (F)"),
 ("OD Room Temps", "OD Room Temps", [("PID_PV_OD_DB_TEMP_STS", "OD Db Temps", 0), ("PID_PV_OD_WB_TEMP_STS", "OD Wb Temps", 1)], "Drybulb Temp (F)", "Wetbulb Temp (F)"),
 ("Amps and Volts", "Amps and Volts", [("UUT_ID_UNIT_AMPS", "Amps (A)", 0), ("UUT_ID_UNIT_VOLTS", "Volts (V)", 1)], "Amps (A)", "Volts (V)"),
 ("Wattage (W)", "Watts", [("UUT_ID_UNIT_WATTS", "Wattage (W)", 0)], "Wattage (W)", None),
 ("ID ESP", "ID ESP", [("PID_PV_ID_TNL_AIRFLOW_PRESS_STS", "ID ESP", 0)], "ID ESP", None),
 ("ID Pressure", "ID Pressure", [(f"UUT_REFRIG_PRESS_100{i}", f"UUT_REFRIG_PRESS_100{i}", 0) for i in (1, 2, 3)], "ID Pressure", None),
 ("Capacity", "Indoor Capacity", [("ID_TOTAL_CAPACITY_REPORT", "Total Capacity Indoor Side", 0),
   ("ID_SENSIBLE_CAPACITY_REPORT", "Sensible Capacity Indoor Side", 0), ("ID_LATENT_CAPACITY_REPORT", "Latent Capacity Indoor Side", 0)], "BTUs", None),
]

def _chart(title, xt, yt):
    c = ScatterChart(); c.title = title; c.scatterStyle = "smoothMarker"
    c.varyColors = False                 # one colour per series, not per point
    c.x_axis.title = xt; c.y_axis.title = yt
    c.x_axis.delete = False; c.y_axis.delete = False
    c.x_axis.tickLblPos = "low"          # keep X labels at the bottom for negative data
    c.width, c.height = 22, 11
    return c

def write_cts_charts(wb, df):
    chs = [(t, ct, sp, y1, y2) for t, ct, sp, y1, y2 in CHARTS if any(s[0] in df.columns for s in sp)]
    cols = ["Time"] + sorted({s[0] for _, _, sp, _, _ in chs for s in sp if s[0] in df.columns})
    data = wb.create_sheet("CTS Chart Data"); data.append(cols)
    for row in df[cols].itertuples(index=False): data.append(list(row))
    n = len(df); x = Reference(data, min_col=1, min_row=2, max_row=n + 1)
    ws = wb.create_sheet("CTS Charts"); anchor = 1
    for _, ctitle, spec, y1, y2 in chs:
        prim = _chart(ctitle, "Time (s)", y1); sec = _chart(ctitle, "Time (s)", y2 or "") if y2 else None
        for chn, name, is_sec in spec:
            if chn not in df.columns: continue
            ref = Reference(data, min_col=cols.index(chn) + 1, min_row=2, max_row=n + 1)
            s = Series(ref, x, title=name); s.marker.symbol = "none"; s.smooth = True
            (sec if (is_sec and sec) else prim).series.append(s)
        if sec and sec.series:
            sec.y_axis.axId = 200; sec.y_axis.crosses = "max"; prim += sec
        if len(prim.series) + (len(sec.series) if sec else 0) == 1:
            prim.legend = None
        ws.add_chart(prim, f"A{anchor}"); anchor += 24

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cts", required=True); ap.add_argument("--agilent")
    ap.add_argument("--mapping", required=True); ap.add_argument("--out", default="AP4_report.xlsx")
    ap.add_argument("--lo", type=float, default=-50); ap.add_argument("--hi", type=float, default=1000)
    ap.add_argument("--include-ambient", action="store_true", help="keep Agilent 1025 (ambient) in max-temp ranking")
    ap.add_argument("--end", default=None, help="stabilization end time hh:mm (default: end of data)")
    a = ap.parse_args()

    cts = parse_cts(a.cts); mp = read_mapping(a.mapping)
    print(f"Mapping: {len(mp)} Agilent channels")
    if not mp:
        raise SystemExit("No channels read from mapping workbook")
    wb = Workbook(); wb.remove(wb.active)

    if a.agilent:
        ag_raw = parse_agilent(a.agilent)
        ti = int(round((ag_raw["Time"].iloc[1] - ag_raw["Time"].iloc[0]).total_seconds()))
        ag = trim_agilent(ag_raw, cts["Date/Time"].iloc[0], cts["Date/Time"].iloc[-1], ti)
        print(f"Agilent interval {ti}s; cropped to {ag['Time'].iloc[0]} .. {ag['Time'].iloc[-1]} ({len(ag)} rows)")

        ag_cols = {c.split(" ", 1)[0]: c for c in ag.columns[1:]}   # "1001" -> "1001 (F)"
        items = []
        for m in mp:
            if m["channel"] == "1025" and not a.include_ambient:
                print("Skipping 1025 (ambient reference); use --include-ambient to keep it")
                continue
            col = ag_cols.get(m["channel"])
            if col is None:
                print(f"WARNING: mapping channel {m['channel']} not found in Agilent data")
                continue
            items.append((m["name"], m["channel"], col, m["limit_c"]))
        write_stats(wb, "Agilent Max Temps", ag, items, a.lo, a.hi, resolve_end(ag, a.end))
    write_cts_charts(wb, cts)
    wb.save(a.out); print("saved", a.out)

if __name__ == "__main__":
    main()
